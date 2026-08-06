#!/usr/bin/env python3
"""ARC internal line-item workbook — 3231 Appalachian Trl (water-only, A/B options).
Rebuild priced on real Xactimate minimum-charge floors (covers sub + ARC margin).
Formula-driven off a JOB ASSUMPTIONS block. No hardcoded totals."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY="16243F"; RED="C41230"; YEL="FFF2CC"; BLUE1="DDE4EE"; GREY="6B7684"; PULLC="FCE4D6"; GRN="1F7A1F"
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Internal Line Items"
thin=Side(style="thin",color="CBD3DC"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def C(cell,v,bold=False,color="1A1A1A",size=10,fill=None,align=None,italic=False,border=False,fmt=None,wrap=False):
    cell.value=v
    cell.font=Font(name="Helvetica",bold=bold,color=color,size=size,italic=italic)
    if fill: cell.fill=PatternFill("solid",fgColor=fill)
    cell.alignment=Alignment(horizontal=align if align else "general",vertical="center",wrap_text=wrap)
    if border: cell.border=box
    if fmt: cell.number_format=fmt
for c,w in {"A":5,"B":9,"C":11,"D":46,"E":6,"F":9,"G":11,"H":12,"I":40}.items(): ws.column_dimensions[c].width=w

r=1
ws.merge_cells(f"A{r}:I{r}"); C(ws[f"A{r}"],"AMERICA'S RESTORATION CONTRACTOR  —  INTERNAL ESTIMATE WORKBOOK",bold=True,color=NAVY,size=13); r+=1
ws.merge_cells(f"A{r}:I{r}"); C(ws[f"A{r}"],"3231 Appalachian Trl, Kingwood, TX 77345  ·  Master Bathroom  ·  Category: WTR (Water — no mold / no containment)  ·  Price List: TXHO8X_JUL26",color=GREY,size=9); r+=1
ws.merge_cells(f"A{r}:I{r}"); C(ws[f"A{r}"],"Retail / cash — price-sensitive, 3 bids.  Option A = minimal patch.  Option B = proper (detach both vanities, full base R&R, full repaint).  Rebuild is SUBBED — priced to cover sub + ARC margin.",color=GREY,size=9,italic=True); r+=2

ws.merge_cells(f"A{r}:I{r}"); C(ws[f"A{r}"],"JOB ASSUMPTIONS  (edit yellow cells — all quantities & margins recompute)",bold=True,color="FFFFFF",fill=NAVY,size=10); r+=1
A={}
def assume(label,val,key,note="",fmt=None):
    global r
    C(ws[f"B{r}"],label,size=9); ws.merge_cells(f"B{r}:C{r}")
    C(ws[f"D{r}"],val,bold=True,fill=YEL,align="center",border=True,fmt=fmt)
    C(ws[f"E{r}"],note,color=GREY,size=8,italic=True); ws.merge_cells(f"E{r}:I{r}")
    A[key]=f"$D${r}"; r+=1
assume("Ceiling height (ft)",8.0,"ceil","CONFIRM — assumed standard 8'")
assume("Water category (1/2/3)",2,"cat","aged stain w/ spotting → treat as Cat 2 (grey)")
assume("Drying days",2,"dry","small wall cavity; drop if meters read dry on arrival")
assume("Site days (Option A)",3,"sdA","drying days + 1 (setup + monitor + takedown)")
assume("Extra site days (Option B)",1,"sdBx","vanity R&R + full repaint adds ~1 day")
assume("Techs on site",1,"tech","")
assume("Monitoring hrs / tech / day",1.75,"mhr","ARC established basis — carries drive time to Kingwood")
assume("A: drywall cut width (ft)",2.0,"acw","")
assume("A: drywall cut height (ft)",2.0,"ach","")
assume("A: affected baseboard (LF)",3,"abase","")
assume("A: antimicrobial surface (SF)",16,"aamb","cut cavity + surround")
assume("B: vanities detach & reset (LF)",6.5,"bvan","his + hers combined face LF — CONFIRM")
assume("B: full baseboard D&R (LF)",28,"bbase","full affected-area perimeter — CONFIRM")
assume("B: full-area wall repaint (SF)",200,"bpaint","paintable grey walls, excl. tile accent — CONFIRM")
assume("Air movers",1,"am","")
assume("Dehumidifiers",1,"dh","")
assume("Rebuild sub cost (% of rebuild billed)",0.68,"subp","what the sub charges ARC — edit to your sub's real number","0%")
assume("Mitigation ARC out-of-pocket ($)",60,"moop","self-performed; materials only (equipment owned)",'"$"#,##0')
assume("Material % of line total (tax basis)",0.40,"matpct","placeholder — true up from Xactimate 'Mat Only'","0%")
assume("Material sales tax rate",0.0825,"taxr","Houston 8.25%, materials only","0.00%")
r+=1
DW=f"({A['acw']}*{A['ach']})"; EQA=f"({A['mhr']}*{A['tech']}*{A['sdA']})"; EQBX=f"({A['mhr']}*{A['tech']}*{A['sdBx']})"

hdr=["#","Cat","Sel","Description","Unit","Qty","Unit $","Extension","Note"]
def header_row():
    global r
    for i,h in enumerate(hdr):
        C(ws.cell(r,i+1),h,bold=True,color="FFFFFF",fill=NAVY,size=9,align=("right" if h in("Qty","Unit $","Extension") else "left"),border=True)
    r+=1
n=[0]; ext={"MIT":[],"A":[],"B":[]}
def line(cat,sel,desc,unit,qty,price,note="",bucket="MIT",pull=False):
    global r
    n[0]+=1
    C(ws.cell(r,1),n[0],size=9,border=True,align="center"); C(ws.cell(r,2),cat,size=9,border=True)
    C(ws.cell(r,3),sel,size=9,border=True); C(ws.cell(r,4),desc,size=9,border=True,wrap=True)
    C(ws.cell(r,5),unit,size=9,border=True,align="center")
    C(ws.cell(r,6),None,size=9,border=True,align="right",fmt="0.00"); ws.cell(r,6).value=f"={qty}"
    if pull:
        C(ws.cell(r,7),"PULL",bold=True,fill=PULLC,size=9,border=True,align="center")
        C(ws.cell(r,8),0,size=9,border=True,align="right",fmt='"$"#,##0.00')
    else:
        C(ws.cell(r,7),price,size=9,border=True,align="right",fmt='"$"#,##0.00')
        C(ws.cell(r,8),None,bold=True,size=9,border=True,align="right",fmt='"$"#,##0.00'); ws.cell(r,8).value=f"=F{r}*G{r}"
    C(ws.cell(r,9),note,color=GREY,size=8,italic=True,border=True,wrap=True)
    ext[bucket].append(f"H{r}"); r+=1
def section(title,color=RED):
    global r; ws.merge_cells(f"A{r}:I{r}"); C(ws[f"A{r}"],title,bold=True,color="FFFFFF",fill=color,size=10); r+=1
def ref(key): return A[key].replace('$','')

section("SECTION 1 — WATER MITIGATION & DRY-OUT  (common to both options)",NAVY); header_row()
line("WTR","DRYW","Tear out wet drywall, cleanup, bag for disposal","SF",DW,1.16,"2'x2' affected cut","MIT")
line("WTR","GRMB","Apply plant-based anti-microbial agent to the surface area","SF",ref('aamb'),0.38,"cavity + surround","MIT")
line("WTR","DRY","Air mover (per 24 hr period) - no monitoring","EA",f"{ref('am')}*{ref('dry')}",26.77,"1 unit x drying days","MIT")
line("WTR","DHM","Dehumidifier (per 24 hr period) - up to 69 ppd","EA",f"{ref('dh')}*{ref('dry')}",60.24,"1 unit x drying days; drop if dry on arrival","MIT")
line("WTR","EQ","Equipment setup, take down, and monitoring (hourly)","HR",EQA,66.75,"1.75 x techs x site days (A). Carries drive time.","MIT")
line("CON","LAB","Content manipulation charge - per hour","HR","1",50.06,"move counter items / work around vanity","MIT")

section("SECTION 2A — REBUILD · OPTION A (minimal patch)  ·  SUBBED",RED); header_row()
line("DRY","MNRP","Drywall repair - minimum charge - labor & material","EA","1",611.16,"cut, hang, tape, float, texture 2x2 patch — real sub floor","A")
line("PNT","MNRP","Painting - minimum charge - labor & material","EA","1",316.08,"prime + paint repaired wall to blend","A")
line("WTR","BASERS","Baseboard - detach & reset","LF",ref('abase'),2.48,"reuse affected base","A")

section("SECTION 2B — REBUILD · OPTION B (proper: detach vanities, full base R&R, full repaint)  ·  SUBBED",RED); header_row()
line("DRY","MNRP","Drywall repair - minimum charge - labor & material","EA","1",611.16,"same 2x2 repair, real sub floor","B")
line("CAB","VANTRS","Vanity with countertop - detach & reset","LF",ref('bvan'),111.58,"both vanities (his + hers), quartz tops","B")
line("WTR","BASERS","Baseboard - detach & reset","LF",ref('bbase'),2.48,"full area","B")
line("PNT","SP2","Seal/prime (1 coat) then paint (2 coats)","SF",ref('bpaint'),1.67,"full repaint of area — exceeds paint min, SF governs","B")
line("PNT","B2","Paint baseboard - two coats","LF",ref('bbase'),1.83,"full area base","B")
line("CLN","FCT","Clean floor - tile","SF","60",0.78,"final clean, larger area","B")
line("WTR","EQ","Equipment setup, take down & monitoring - added day","HR",EQBX,66.75,"extra monitoring/labor day for Option B scope","B")
r+=1

mit="+".join(ext["MIT"]); aa="+".join(ext["A"]); bb="+".join(ext["B"])
def total_block(title,sub):
    global r
    C(ws.cell(r,4),title,bold=True,color=NAVY,size=10,align="right")
    C(ws.cell(r,8),None,bold=True,color=NAVY,size=10,border=True,align="right",fmt='"$"#,##0.00'); sr=r; ws.cell(r,8).value=f"={sub}"; r+=1
    C(ws.cell(r,4),"Est. material sales tax (materials only)",size=9,align="right",color=GREY)
    C(ws.cell(r,8),None,size=9,border=True,align="right",fmt='"$"#,##0.00',color=GREY); tr=r; ws.cell(r,8).value=f"=ROUND(H{sr}*{A['matpct']}*{A['taxr']},2)"; r+=1
    C(ws.cell(r,4),title.replace("Subtotal","TOTAL"),bold=True,color="FFFFFF",fill=NAVY,size=10,align="right")
    for cc in (5,6,7): ws.cell(r,cc).fill=PatternFill("solid",fgColor=NAVY)
    C(ws.cell(r,8),None,bold=True,color="FFFFFF",fill=NAVY,size=11,align="right",fmt='"$"#,##0.00'); tot=r; ws.cell(r,8).value=f"=H{sr}+H{tr}"; r+=1
    return sr,tot
section("OPTION A — TOTAL (Mitigation + Rebuild A)",NAVY); Asub,Atot=total_block("Option A Subtotal",f"{mit}+{aa}"); r+=1
section("OPTION B — TOTAL (Mitigation + Rebuild B)",NAVY); Bsub,Btot=total_block("Option B Subtotal",f"{mit}+{bb}"); r+=2

# ---------- MARGIN: what ARC keeps after paying the rebuild sub ----------
section("MARGIN — WHAT ARC KEEPS AFTER PAYING THE REBUILD SUB",NAVY)
C(ws.cell(r,2),"Mitigation is self-performed (ARC keeps it, less materials). Rebuild is subbed. ARC keep = (mitigation - materials) + (rebuild billed - sub cost).",color=GREY,size=9,italic=True,wrap=True); ws.merge_cells(f"B{r}:I{r}"); ws.row_dimensions[r].height=26; r+=1
mitsum=f"({mit})"
def margin_block(label,rebuild_sum,subtot_row):
    global r
    C(ws.cell(r,2),label,bold=True,color=RED,size=10); ws.merge_cells(f"B{r}:I{r}"); r+=1
    def row(lbl,formula,fmt='"$"#,##0',color="1A1A1A",bold=False):
        global r
        C(ws.cell(r,3),lbl,size=9,color="333333"); ws.merge_cells(f"C{r}:F{r}")
        C(ws.cell(r,8),None,size=9,bold=bold,align="right",fmt=fmt,color=color); ws.cell(r,8).value=f"={formula}"
        rr=r; r+=1; return rr
    mb=row("Mitigation billed (self-perform)",mitsum)
    mo=row("  less mitigation materials (OOP)",f"-{A['moop']}",color=GREY)
    rb=row("Rebuild billed",f"({rebuild_sum})")
    sc=row("  less rebuild sub cost",f"-({rebuild_sum})*{A['subp']}",color=GREY)
    C(ws.cell(r,3),"  = rebuild margin ARC keeps",size=9,color="333333"); ws.merge_cells(f"C{r}:F{r}")
    C(ws.cell(r,8),None,size=9,align="right",fmt='"$"#,##0',color=GRN); ws.cell(r,8).value=f"=({rebuild_sum})-({rebuild_sum})*{A['subp']}"
    C(ws.cell(r,9),None,size=9,align="left",fmt="0%",color=GRN,italic=True); ws.cell(r,9).value=f"=1-{A['subp']}"; r+=1
    C(ws.cell(r,3),"ARC TOTAL CONTRIBUTION",bold=True,size=10,color=NAVY); ws.merge_cells(f"C{r}:F{r}")
    C(ws.cell(r,8),None,bold=True,size=10,align="right",fmt='"$"#,##0',color=GRN,border=True)
    ws.cell(r,8).value=f"=(H{mb}+H{mo})+(({rebuild_sum})-({rebuild_sum})*{A['subp']})"
    C(ws.cell(r,9),None,bold=True,size=9,align="left",fmt="0%",color=GRN,italic=True)
    ws.cell(r,9).value=f"=H{r}/H{subtot_row}"; C(ws.cell(r,9),ws.cell(r,9).value); r+=2
margin_block("OPTION A",aa,Asub)
margin_block("OPTION B",bb,Bsub)
C(ws.cell(r,2),"Suggested sell (round):",bold=True,color=RED,size=10); ws.merge_cells(f"B{r}:D{r}")
C(ws.cell(r,5),"Option A ≈ $1,500–$1,600   ·   Option B ≈ $2,500–$2,650",bold=True,color=NAVY,size=10); ws.merge_cells(f"E{r}:I{r}"); r+=1
C(ws.cell(r,2),"Biggest margin lever (primer §6): self-perform the rebuild instead of subbing — turns the sub cost into ~material only and drops it straight to ARC. Make the next hire rebuild-capable.",color=GREY,size=8,italic=True,wrap=True); ws.merge_cells(f"B{r}:I{r}"); ws.row_dimensions[r].height=26; r+=2

section("NOTES / STRATEGY / DISCLOSURES",NAVY)
notes=[
 "CATEGORY: Built as WTR (water) per direction — no mold remediation, no containment, no HMR rates. A small dark spot is visible at the base of the affected wall; if opening reveals mold >25 contiguous SF, TX TDLR requires stop-work and conversion to a licensed mold path (separate scope).",
 "REBUILD PRICING (corrected): rebuild is subbed, so it is priced on the real Xactimate minimum-charge floors (DRY MNRP $611 + PNT MNRP $316, both labor+material) — this is what a sub actually charges to show up for a small repair, NOT a maximized build. Covers the sub AND leaves ARC margin. Earlier draft priced bare per-SF production and underpriced the rebuild — that is the leak the sub has been eating.",
 "REBUILD MARGIN: at a 68% sub-cost assumption the rebuild carries ~32% ARC margin — inside the 30-40% target. Edit 'Rebuild sub cost %' to your sub's real quote; self-performing pushes margin much higher.",
 "MITIGATION-to-REBUILD RATIO: Option A runs ~1 : 1.6 (mit ~$585 / rebuild ~$935), in line with the 1 : 2 rule other contractors quoted. Option B is heavier on rebuild because of the two-vanity detach & reset.",
 "MONITORING (EQ): WTR rate $66.75/hr used (NOT the HMR $78.34). This line carries the drive time to Kingwood — do not cut it to soften the total.",
 "MOISTURE SOURCE: The affected wall backs the his/hers vanity plumbing. Probable source is a vanity supply/drain or caulk failure, NOT in this scope. Disclose in writing that a plumber should verify/repair the source before close-up or it can recur. This disclosure protects ARC.",
 "INSULATION: Not included — assumed interior partition. If wet insulation is found, add WTR INS ($0.86/SF) by change order.",
 "BASEBOARD: Detach & reset (reuse). If a piece is swollen, replace only that piece — FCW LAMB $7.62/LF is the only base-install line in the current export; verify base profile in Xactimate before quoting a full replace.",
 "TAX: Materials-only 8.25% on a 40% material placeholder — an ESTIMATE. Texas labor for residential restoration is not taxable. True up from Xactimate 'Mat Only' before send.",
 "CUSTOMER wants a cost breakdown: proposal shows price by scope section (mitigation / rebuild) per option — enough to satisfy the request without itemizing machines or day counts. No mold testing included (TDLR conflict-of-interest).",
]
for i,t in enumerate(notes,1):
    C(ws.cell(r,2),f"{i}.",bold=True,size=8,color=GREY)
    C(ws.cell(r,3),t,size=8,color="333333",wrap=True); ws.merge_cells(f"C{r}:I{r}")
    ws.row_dimensions[r].height=14+(len(t)//95)*11; r+=1
ws.freeze_panes="A6"
out="/tmp/claude-0/-home-user-RXS/0ad0031b-0623-52ca-a8ef-bf297e37d330/scratchpad/ARC_Estimate_3231_Appalachian_INTERNAL.xlsx"
wb.save(out); print("saved",out); print("Asub",Asub,"Atot",Atot,"Bsub",Bsub,"Btot",Btot)
